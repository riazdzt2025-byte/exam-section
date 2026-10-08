from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Activity, AppSetting, Subject, Todo
from .utils import local_today

User = get_user_model()


class AuthFlowTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("admin1", password="adminpass1", first_name="এডমিন", is_staff=True)

    def test_login_required_for_home(self):
        r = self.client.get(reverse("home"))
        self.assertRedirects(r, f"{reverse('login')}?next=/")

    def test_no_self_registration_url(self):
        self.assertEqual(self.client.get("/register/").status_code, 404)
        self.assertEqual(self.client.get("/signup/").status_code, 404)

    def test_admin_creates_member_and_member_can_login(self):
        self.client.login(username="admin1", password="adminpass1")
        r = self.client.post(reverse("members"), {
            "username": "saymun", "display_name": "সাইমুন হোসেন", "password": "secret12",
        })
        self.assertRedirects(r, reverse("members"))
        member = User.objects.get(username="saymun")
        self.assertFalse(member.is_staff)
        self.assertNotEqual(member.password, "secret12")  # হ্যাশ করা
        self.client.logout()
        self.assertTrue(self.client.login(username="saymun", password="secret12"))

    def test_admin_resets_member_password(self):
        member = User.objects.create_user("habib", password="oldpass12")
        self.client.login(username="admin1", password="adminpass1")
        self.client.post(reverse("member_password", args=[member.pk]), {"password": "newpass12"})
        self.client.logout()
        self.assertFalse(self.client.login(username="habib", password="oldpass12"))
        self.assertTrue(self.client.login(username="habib", password="newpass12"))

    def test_short_password_rejected(self):
        member = User.objects.create_user("emon", password="oldpass12")
        self.client.login(username="admin1", password="adminpass1")
        self.client.post(reverse("member_password", args=[member.pk]), {"password": "123"})
        member.refresh_from_db()
        self.assertTrue(member.check_password("oldpass12"))

    def test_member_cannot_manage_members(self):
        User.objects.create_user("sorab", password="memberpass1")
        self.client.login(username="sorab", password="memberpass1")
        self.assertEqual(self.client.get(reverse("members")).status_code, 403)
        victim = User.objects.create_user("x", password="oldpass12")
        r = self.client.post(reverse("member_password", args=[victim.pk]), {"password": "hacked123"})
        self.assertEqual(r.status_code, 403)
        victim.refresh_from_db()
        self.assertTrue(victim.check_password("oldpass12"))

    def test_admin_cannot_deactivate_self(self):
        self.client.login(username="admin1", password="adminpass1")
        self.client.post(reverse("member_toggle", args=[self.admin.pk]))
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.is_active)

    def test_deactivated_member_cannot_login(self):
        m = User.objects.create_user("off", password="memberpass1", is_active=False)
        self.assertFalse(self.client.login(username="off", password="memberpass1"))


class TaskPageTestBase(TestCase):
    """সদস্যদের কাজ পেজ ও পারমিশনের পরীক্ষার ভিত্তি।"""

    def setUp(self):
        self.admin = User.objects.create_user("admin1", password="adminpass1", first_name="এডমিন", is_staff=True)
        self.saimun = User.objects.create_user("saimun", password="memberpass1", first_name="সাইমুন হোসেন")
        self.emon = User.objects.create_user("emon", password="memberpass1", first_name="ইমন হোসেন")
        Subject.objects.create(owner=self.saimun, name="বাংলা ১ম পত্র", submitted=True)

    def login_saimun(self):
        self.assertTrue(self.client.login(username="saimun", password="memberpass1"))

    def login_admin(self):
        self.assertTrue(self.client.login(username="admin1", password="adminpass1"))


class TasksPageTests(TaskPageTestBase):
    def test_tasks_requires_login(self):
        r = self.client.get(reverse("tasks"))
        self.assertRedirects(r, f"{reverse('login')}?next={reverse('tasks')}")

    def test_member_sees_edit_controls_only_on_own_card(self):
        self.login_saimun()
        html = self.client.get(reverse("tasks")).content.decode()
        self.assertIn(f'action="{reverse("subject_add", args=[self.saimun.pk])}"', html)
        self.assertNotIn(f'action="{reverse("subject_add", args=[self.emon.pk])}"', html)

    def test_member_adds_subject_to_self(self):
        self.login_saimun()
        r = self.client.post(reverse("subject_add", args=[self.saimun.pk]), {"name": "ইংরেজি", "note": "৩য় সেম"})
        self.assertRedirects(r, reverse("tasks"))
        s = Subject.objects.get(owner=self.saimun, name="ইংরেজি")
        self.assertEqual(s.note, "৩য় সেম")
        self.assertTrue(Activity.objects.filter(kind=Activity.SUBJECT_ADD).exists())

    def test_member_cannot_add_subject_to_other(self):
        self.login_saimun()
        r = self.client.post(reverse("subject_add", args=[self.emon.pk]), {"name": "হ্যাক"})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(Subject.objects.filter(owner=self.emon).exists())

    def test_duplicate_subject_rejected(self):
        self.login_saimun()
        self.client.post(reverse("subject_add", args=[self.saimun.pk]), {"name": "বাংলা ১ম পত্র"})
        self.assertEqual(Subject.objects.filter(owner=self.saimun, name="বাংলা ১ম পত্র").count(), 1)

    def test_admin_can_add_subject_to_anyone(self):
        self.login_admin()
        self.client.post(reverse("subject_add", args=[self.emon.pk]), {"name": "গণিত"})
        self.assertTrue(Subject.objects.filter(owner=self.emon, name="গণিত").exists())

    def test_member_toggles_own_subject(self):
        self.login_saimun()
        s = Subject.objects.get(owner=self.saimun)
        self.client.post(reverse("subject_toggle", args=[s.pk, "typed"]))
        s.refresh_from_db()
        self.assertTrue(s.typed)
        self.assertTrue(Activity.objects.filter(kind=Activity.SUBJECT_STATUS).exists())

    def test_member_cannot_toggle_other_subject(self):
        self.login_saimun()
        s = Subject.objects.create(owner=self.emon, name="সুরক্ষিত")
        r = self.client.post(reverse("subject_toggle", args=[s.pk, "submitted"]))
        self.assertEqual(r.status_code, 403)
        s.refresh_from_db()
        self.assertFalse(s.submitted)

    def test_invalid_toggle_field_rejected(self):
        self.login_admin()
        s = Subject.objects.get(owner=self.saimun)
        r = self.client.post(reverse("subject_toggle", args=[s.pk, "is_superuser"]))
        self.assertEqual(r.status_code, 403)

    def test_member_cannot_delete_other_subject(self):
        self.login_saimun()
        s = Subject.objects.create(owner=self.emon, name="মুছবে না")
        r = self.client.post(reverse("subject_delete", args=[s.pk]))
        self.assertEqual(r.status_code, 403)
        self.assertTrue(Subject.objects.filter(pk=s.pk).exists())

    def test_admin_deletes_any_subject_and_logs(self):
        self.login_admin()
        s = Subject.objects.create(owner=self.emon, name="মুছবে")
        self.client.post(reverse("subject_delete", args=[s.pk]))
        self.assertFalse(Subject.objects.filter(pk=s.pk).exists())
        self.assertTrue(Activity.objects.filter(kind=Activity.SUBJECT_DELETE).exists())

    def test_filter_preserved_after_toggle(self):
        self.login_saimun()
        s = Subject.objects.get(owner=self.saimun)
        r = self.client.post(reverse("subject_toggle", args=[s.pk, "typed"]), {"filter": "pending"})
        self.assertRedirects(r, reverse("tasks") + "?filter=pending")

    def test_filter_pending_and_done(self):
        Subject.objects.create(
            owner=self.saimun, name="সম্পন্ন বিষয়", submitted=True, typed=True, photocopied=True
        )
        self.login_saimun()
        html = self.client.get(reverse("tasks") + "?filter=done").content.decode()
        self.assertIn("<td>সম্পন্ন বিষয়</td>", html)
        self.assertNotIn("<td>বাংলা ১ম পত্র</td>", html)
        html = self.client.get(reverse("tasks") + "?filter=pending").content.decode()
        self.assertIn("<td>বাংলা ১ম পত্র</td>", html)
        self.assertNotIn("<td>সম্পন্ন বিষয়</td>", html)


class TodoTests(TaskPageTestBase):
    def test_member_adds_todo_to_self_only(self):
        self.login_saimun()
        self.client.post(reverse("todo_add", args=[self.saimun.pk]), {"text": "বন্ড করা"})
        self.assertTrue(Todo.objects.filter(owner=self.saimun, text="বন্ড করা").exists())
        r = self.client.post(reverse("todo_add", args=[self.emon.pk]), {"text": "না"})
        self.assertEqual(r.status_code, 403)

    def test_todo_toggle_and_delete(self):
        self.login_saimun()
        t = Todo.objects.create(owner=self.saimun, text="কাজ ১")
        self.client.post(reverse("todo_toggle", args=[t.pk]))
        t.refresh_from_db()
        self.assertTrue(t.done)
        self.client.post(reverse("todo_delete", args=[t.pk]))
        self.assertFalse(Todo.objects.filter(pk=t.pk).exists())

    def test_member_cannot_touch_other_todo(self):
        t = Todo.objects.create(owner=self.emon, text="সুরক্ষিত")
        self.login_saimun()
        self.assertEqual(self.client.post(reverse("todo_toggle", args=[t.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("todo_delete", args=[t.pk])).status_code, 403)
        self.assertTrue(Todo.objects.filter(pk=t.pk).exists())


class DashboardTests(TaskPageTestBase):
    def test_kpi_counts_and_report(self):
        Subject.objects.create(
            owner=self.emon, name="গণিত", submitted=True, typed=True, photocopied=True
        )
        self.login_admin()
        r = self.client.get(reverse("home"))
        totals = r.context["totals"]
        self.assertEqual(totals["total"], 2)
        self.assertEqual(totals["finished"], 1)
        self.assertEqual(totals["percent"], 50)
        self.assertContains(r, "মোট বিষয়")
        self.assertIn("মোট বিষয়: ২", r.context["report"])  # বাংলা অঙ্ক
        self.assertContains(r, "copyReport")  # কপি বাটন
        self.assertContains(r, "window.print()")  # প্রিন্ট বাটন


class DailyReportTests(TaskPageTestBase):
    def test_requires_login(self):
        r = self.client.get(reverse("daily_report"))
        self.assertRedirects(r, f"{reverse('login')}?next={reverse('daily_report')}")

    def test_today_activity_visible(self):
        self.login_admin()
        self.client.post(reverse("subject_add", args=[self.emon.pk]), {"name": "আজকের বিষয়"})
        r = self.client.get(reverse("daily_report"))
        self.assertContains(r, "আজকের বিষয়")
        self.assertEqual(r.context["totals"]["subject_add"], 1)

    def test_old_date_shows_nothing(self):
        self.login_admin()
        self.client.post(reverse("subject_add", args=[self.emon.pk]), {"name": "আজকের বিষয়"})
        r = self.client.get(reverse("daily_report"), {"date": "2000-01-01"})
        self.assertEqual(r.context["activities"].count(), 0)
        self.assertContains(r, "কোনো কার্যক্রম পাওয়া যায়নি")

    def test_invalid_date_falls_back_to_today(self):
        self.login_admin()
        r = self.client.get(reverse("daily_report"), {"date": "আবুল"})
        self.assertEqual(r.context["day"], local_today())


class ExcelTests(TaskPageTestBase):
    def _make_import_file(self, rows, todos=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Subjects"
        ws.append(["username", "subject", "note", "submitted", "typed", "photocopied"])
        for row in rows:
            ws.append(row)
        if todos is not None:
            ws2 = wb.create_sheet("Todos")
            ws2.append(["username", "text", "done"])
            for row in todos:
                ws2.append(row)
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        out.name = "import.xlsx"
        return out

    def test_export_returns_valid_xlsx(self):
        Todo.objects.create(owner=self.saimun, text="কাজ")
        self.login_saimun()
        r = self.client.get(reverse("excel_export"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        wb = load_workbook(BytesIO(r.content))
        self.assertIn("Subjects", wb.sheetnames)
        self.assertIn("Todos", wb.sheetnames)
        data = list(wb["Subjects"].iter_rows(values_only=True))
        self.assertEqual(data[1][0], "saimun")
        self.assertEqual(data[1][2], "বাংলা ১ম পত্র")
        self.assertEqual(data[1][4], "yes")  # submitted

    def test_template_download(self):
        self.login_admin()
        r = self.client.get(reverse("excel_template"))
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content))
        self.assertIn("নির্দেশনা", wb.sheetnames)

    def test_import_creates_updates_and_skips(self):
        self.login_admin()
        f = self._make_import_file(
            [
                ["emon", "গণিত", "২য় বর্ষ", "yes", "no", "no"],      # নতুন
                ["saimun", "বাংলা ১ম পত্র", "", "yes", "yes", None],  # হালনাগাদ (টাইপ)
                ["ghost", "ভুত বিষয়", "", "no", "no", "no"],          # অজানা ইউজার → বাদ
                [None, "", "", "", "", ""],                            # খালি সারি → বাদ
            ],
            todos=[["emon", "প্রশ্ন বন্ড", "no"]],
        )
        self.client.post(reverse("excel_import"), {"file": f})
        new = Subject.objects.get(owner=self.emon, name="গণিত")
        self.assertTrue(new.submitted)
        saimun_s = Subject.objects.get(owner=self.saimun, name="বাংলা ১ম পত্র")
        self.assertTrue(saimun_s.typed)  # হালনাগাদ হয়েছে
        self.assertTrue(saimun_s.submitted)  # আগের মান অক্ষত
        self.assertFalse(Subject.objects.filter(name="ভুত বিষয়").exists())
        self.assertTrue(Todo.objects.filter(owner=self.emon, text="প্রশ্ন বন্ড").exists())
        # একই টেক্সটের টুডু দ্বিতীয়বারে নতুন করে তৈরি হয় না
        f2 = self._make_import_file([], todos=[["emon", "প্রশ্ন বন্ড", "yes"]])
        self.client.post(reverse("excel_import"), {"file": f2})
        self.assertEqual(Todo.objects.filter(owner=self.emon, text="প্রশ্ন বন্ড").count(), 1)

    def test_import_admin_only(self):
        self.login_saimun()
        f = self._make_import_file([["emon", "গণিত", "", "no", "no", "no"]])
        r = self.client.post(reverse("excel_import"), {"file": f})
        self.assertEqual(r.status_code, 403)

    def test_import_rejects_non_xlsx(self):
        self.login_admin()
        out = BytesIO(b"not an excel file")
        out.name = "bad.txt"
        self.client.post(reverse("excel_import"), {"file": out})
        self.assertFalse(Subject.objects.filter(name__contains="bad").exists())


class SettingsTests(TaskPageTestBase):
    def test_admin_updates_exam_title(self):
        self.login_admin()
        self.client.post(reverse("settings_page"), {"exam_title": "বার্ষিক পরীক্ষা ২০২৬"})
        self.assertEqual(AppSetting.exam_title(), "বার্ষিক পরীক্ষা ২০২৬")
        self.assertIn("বার্ষিক পরীক্ষা ২০২৬", self.client.get(reverse("home")).context["report"])

    def test_member_cannot_open_settings(self):
        self.login_saimun()
        self.assertEqual(self.client.get(reverse("settings_page")).status_code, 403)

    def test_blank_title_rejected(self):
        self.login_admin()
        AppSetting.set(AppSetting.EXAM_TITLE, "আগের শিরোনাম")
        self.client.post(reverse("settings_page"), {"exam_title": "   "})
        self.assertEqual(AppSetting.exam_title(), "আগের শিরোনাম")


class ActivityLogTests(TaskPageTestBase):
    def test_old_activity_shown_on_its_own_day(self):
        self.login_admin()
        self.client.post(reverse("subject_add", args=[self.emon.pk]), {"name": "এখনকার কাজ"})
        act = Activity.objects.latest("id")
        Activity.objects.filter(pk=act.pk).update(created=act.created - timedelta(days=3))
        # আজকের প্রতিবেদনে আর নেই, ৩ দিন আগের তারিখে আছে
        r_today = self.client.get(reverse("daily_report"))
        self.assertEqual(r_today.context["activities"].count(), 0)
        gone_day = (local_today() - timedelta(days=3)).isoformat()
        self.assertContains(self.client.get(reverse("daily_report"), {"date": gone_day}), "এখনকার কাজ")
