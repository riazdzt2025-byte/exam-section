from datetime import datetime
from functools import wraps
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import PasswordChangeView
from django.contrib.messages.views import SuccessMessageMixin
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Max, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.views.decorators.http import require_POST

from .forms import (
    ExcelImportForm,
    MemberCreateForm,
    SetMemberPasswordForm,
    SettingsForm,
    SubjectForm,
    TodoForm,
)
from .models import Activity, AppSetting, Subject, Todo
from .utils import day_bounds, local_today, parse_iso_date, to_bn

User = get_user_model()

STATUS_FIELDS = ("submitted", "typed", "photocopied")
STATUS_LABELS = {
    "submitted": "জমা পড়েছে",
    "typed": "টাইপ হয়েছে",
    "photocopied": "ফটোকপি হয়েছে",
}
FILTERS = ("all", "pending", "done")


# ---------------------------------------------------------------- সাধারণ সহায়ক


def admin_required(view):
    """শুধু এডমিন (is_staff) ইউজার ঢুকতে পারবে।"""

    @login_required
    @wraps(view)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_staff:
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapper


def can_edit(user, owner):
    """নিজের ডাটা নিজে, আর সবকিছু এডমিন।"""
    return user.is_staff or user.pk == owner.pk


def _require_edit(request, owner):
    if not can_edit(request.user, owner):
        raise PermissionDenied


def log_activity(actor, owner, kind, text):
    Activity.objects.create(actor=actor, owner=owner, kind=kind, text=text)


def display_name(user):
    return user.first_name or user.username


def tasks_url(flt="all"):
    url = reverse("tasks")
    if flt in FILTERS and flt != "all":
        url += f"?filter={flt}"
    return url


# ---------------------------------------------------------------- ড্যাশবোর্ড


@login_required
def home(request):
    done = Q(subjects__submitted=True, subjects__typed=True, subjects__photocopied=True)
    members = (
        User.objects.filter(is_active=True)
        .annotate(
            total=Count("subjects"),
            finished=Count("subjects", filter=done),
            submitted_n=Count("subjects", filter=Q(subjects__submitted=True)),
            typed_n=Count("subjects", filter=Q(subjects__typed=True)),
            photocopied_n=Count("subjects", filter=Q(subjects__photocopied=True)),
        )
        .order_by("first_name", "username")
    )
    all_subjects = Subject.objects.all()
    totals = {
        "total": all_subjects.count(),
        "submitted": all_subjects.filter(submitted=True).count(),
        "typed": all_subjects.filter(typed=True).count(),
        "photocopied": all_subjects.filter(photocopied=True).count(),
        "finished": all_subjects.filter(submitted=True, typed=True, photocopied=True).count(),
    }
    totals["percent"] = round(totals["finished"] * 100 / totals["total"]) if totals["total"] else 0
    report = build_progress_report(members, totals)
    return render(
        request,
        "core/home.html",
        {
            "members": members,
            "totals": totals,
            "report": report,
            "exam_title": AppSetting.exam_title(),
            "today": local_today(),
        },
    )


def build_progress_report(members, totals):
    """ড্যাশবোর্ডের 'সাংখ্যিক প্রতিবেদন' — কপি করার জন্য সাদা টেক্সট।"""
    lines = [
        AppSetting.exam_title(),
        f"তারিখ: {to_bn(local_today().strftime('%d-%m-%Y'))}",
        "-" * 28,
        f"মোট বিষয়: {to_bn(totals['total'])}",
        f"জমা পড়েছে: {to_bn(totals['submitted'])}",
        f"টাইপ হয়েছে: {to_bn(totals['typed'])}",
        f"ফটোকপি হয়েছে: {to_bn(totals['photocopied'])}",
        f"সব ধাপ শেষ: {to_bn(totals['finished'])} ({to_bn(totals['percent'])}%)",
        "-" * 28,
        "সদস্যভিত্তিক সারাংশ:",
    ]
    for m in members:
        lines.append(
            f"{display_name(m)} — বিষয় {to_bn(m.total)} | জমা {to_bn(m.submitted_n)}"
            f" | টাইপ {to_bn(m.typed_n)} | ফটোকপি {to_bn(m.photocopied_n)}"
            f" | শেষ {to_bn(m.finished)}/{to_bn(m.total)}"
        )
    return "\n".join(lines)


# ---------------------------------------------------------------- সদস্য ব্যবস্থাপনা


@admin_required
def members(request):
    if request.method == "POST":
        form = MemberCreateForm(request.POST)
        if form.is_valid():
            user = form.save()
            messages.success(request, f"সদস্য যোগ হয়েছে: {user.first_name}")
            return redirect("members")
    else:
        form = MemberCreateForm()
    users = User.objects.order_by("-is_active", "first_name", "username")
    return render(request, "core/members.html", {"form": form, "users": users})


@admin_required
@require_POST
def member_password(request, pk):
    user = get_object_or_404(User, pk=pk)
    form = SetMemberPasswordForm(request.POST)
    if form.is_valid():
        user.set_password(form.cleaned_data["password"])
        user.save(update_fields=["password"])
        messages.success(request, f"{user.first_name or user.username}-এর নতুন পাসওয়ার্ড সেট হয়েছে।")
    else:
        messages.error(request, "পাসওয়ার্ড কমপক্ষে ৬ অক্ষরের হতে হবে।")
    return redirect("members")


@admin_required
@require_POST
def member_toggle(request, pk):
    user = get_object_or_404(User, pk=pk)
    if user == request.user:
        messages.error(request, "নিজের অ্যাকাউন্ট বন্ধ করা যাবে না।")
    else:
        user.is_active = not user.is_active
        user.save(update_fields=["is_active"])
        state = "চালু" if user.is_active else "বন্ধ"
        messages.success(request, f"{user.first_name or user.username}-এর অ্যাকাউন্ট {state} হয়েছে।")
    return redirect("members")


class MyPasswordChangeView(SuccessMessageMixin, PasswordChangeView):
    template_name = "core/password_change.html"
    success_url = reverse_lazy("home")
    success_message = "আপনার পাসওয়ার্ড বদলানো হয়েছে।"


# ---------------------------------------------------------------- সদস্যদের কাজ (টাস্ক)


@login_required
def tasks(request):
    flt = request.GET.get("filter", "all")
    if flt not in FILTERS:
        flt = "all"
    cards = []
    for member in User.objects.filter(is_active=True).order_by("first_name", "username"):
        subjects = list(member.subjects.all())
        if flt == "pending":
            subjects = [s for s in subjects if not s.is_done]
        elif flt == "done":
            subjects = [s for s in subjects if s.is_done]
        cards.append(
            {
                "member": member,
                "name": display_name(member),
                "can_edit": can_edit(request.user, member),
                "subjects": subjects,
                "todos": list(member.todos.all()),
                "total": member.subjects.count(),
                "finished": member.subjects.filter(
                    submitted=True, typed=True, photocopied=True
                ).count(),
            }
        )
    return render(
        request,
        "core/tasks.html",
        {"cards": cards, "filter": flt, "fields": STATUS_FIELDS},
    )


@login_required
@require_POST
def subject_add(request, owner_pk):
    owner = get_object_or_404(User, pk=owner_pk, is_active=True)
    _require_edit(request, owner)
    form = SubjectForm(request.POST)
    flt = request.POST.get("filter", "all")
    if form.is_valid():
        name = form.cleaned_data["name"]
        note = form.cleaned_data["note"].strip()
        if owner.subjects.filter(name__iexact=name).exists():
            messages.error(request, f"এই নামে বিষয় আগে থেকেই আছে: {name}")
        else:
            order = (owner.subjects.aggregate(m=Max("order"))["m"] or 0) + 1
            Subject.objects.create(owner=owner, name=name, note=note, order=order)
            log_activity(
                request.user,
                owner,
                Activity.SUBJECT_ADD,
                f"নতুন বিষয় যোগ হয়েছে: {name} — {display_name(owner)}",
            )
            messages.success(request, f"বিষয় যোগ হয়েছে: {name}")
    else:
        messages.error(request, "বিষয়ের নাম লিখুন।")
    return redirect(tasks_url(flt))


@login_required
@require_POST
def subject_toggle(request, pk, field):
    if field not in STATUS_FIELDS:
        raise PermissionDenied
    subject = get_object_or_404(Subject, pk=pk)
    _require_edit(request, subject.owner)
    flt = request.POST.get("filter", "all")
    new_value = not getattr(subject, field)
    setattr(subject, field, new_value)
    subject.save(update_fields=[field])
    state = "হয়েছে" if new_value else "ফিরিয়ে নেওয়া হয়েছে"
    log_activity(
        request.user,
        subject.owner,
        Activity.SUBJECT_STATUS,
        f"{subject.name} ({display_name(subject.owner)}): {STATUS_LABELS[field]} {state}",
    )
    return redirect(tasks_url(flt))


@login_required
@require_POST
def subject_delete(request, pk):
    subject = get_object_or_404(Subject, pk=pk)
    _require_edit(request, subject.owner)
    flt = request.POST.get("filter", "all")
    log_activity(
        request.user,
        subject.owner,
        Activity.SUBJECT_DELETE,
        f"বিষয় মুছে ফেলা হয়েছে: {subject.name} — {display_name(subject.owner)}",
    )
    subject.delete()
    messages.success(request, "বিষয় মুছে ফেলা হয়েছে।")
    return redirect(tasks_url(flt))


@login_required
@require_POST
def todo_add(request, owner_pk):
    owner = get_object_or_404(User, pk=owner_pk, is_active=True)
    _require_edit(request, owner)
    form = TodoForm(request.POST)
    flt = request.POST.get("filter", "all")
    if form.is_valid():
        text = form.cleaned_data["text"]
        Todo.objects.create(owner=owner, text=text)
        log_activity(
            request.user, owner, Activity.TODO_ADD, f"করণীয় যোগ: {text} — {display_name(owner)}"
        )
    return redirect(tasks_url(flt))


@login_required
@require_POST
def todo_toggle(request, pk):
    todo = get_object_or_404(Todo, pk=pk)
    _require_edit(request, todo.owner)
    flt = request.POST.get("filter", "all")
    todo.done = not todo.done
    todo.save(update_fields=["done"])
    state = "সম্পন্ন" if todo.done else "আবার খোলা"
    log_activity(
        request.user,
        todo.owner,
        Activity.TODO_TOGGLE,
        f"করণীয় {state}: {todo.text} — {display_name(todo.owner)}",
    )
    return redirect(tasks_url(flt))


@login_required
@require_POST
def todo_delete(request, pk):
    todo = get_object_or_404(Todo, pk=pk)
    _require_edit(request, todo.owner)
    flt = request.POST.get("filter", "all")
    log_activity(
        request.user,
        todo.owner,
        Activity.TODO_DELETE,
        f"করণীয় মুছে ফেলা হয়েছে: {todo.text} — {display_name(todo.owner)}",
    )
    todo.delete()
    return redirect(tasks_url(flt))


# ---------------------------------------------------------------- দৈনিক প্রতিবেদন


@login_required
def daily_report(request):
    day = parse_iso_date(request.GET.get("date", ""))
    start, end = day_bounds(day)
    activities = (
        Activity.objects.filter(created__gte=start, created__lt=end)
        .select_related("owner", "actor")
        .order_by("created", "id")
    )
    totals = {
        "total": activities.count(),
        "subject_add": activities.filter(kind=Activity.SUBJECT_ADD).count(),
        "subject_status": activities.filter(kind=Activity.SUBJECT_STATUS).count(),
        "subject_delete": activities.filter(kind=Activity.SUBJECT_DELETE).count(),
        "todo": activities.filter(kind__in=[Activity.TODO_ADD, Activity.TODO_TOGGLE]).count(),
    }
    lines = [
        f"{AppSetting.exam_title()} — দৈনিক প্রতিবেদন",
        f"তারিখ: {to_bn(day.strftime('%d-%m-%Y'))}",
        "-" * 28,
        f"মোট কার্যক্রম: {to_bn(totals['total'])}",
        f"নতুন বিষয়: {to_bn(totals['subject_add'])}",
        f"অবস্থা পরিবর্তন: {to_bn(totals['subject_status'])}",
        f"বিষয় মুছে ফেলা: {to_bn(totals['subject_delete'])}",
        f"করণীয় হালনাগাদ: {to_bn(totals['todo'])}",
        "-" * 28,
        "বিবরণ:",
    ]
    for a in activities:
        lines.append(f"{to_bn(a.created.astimezone().strftime('%I:%M %p'))} — {a.text}")
    return render(
        request,
        "core/daily_report.html",
        {
            "day": day,
            "day_iso": day.isoformat(),
            "activities": activities,
            "totals": totals,
            "text_report": "\n".join(lines),
            "is_today": day == local_today(),
            "exam_title": AppSetting.exam_title(),
        },
    )


# ---------------------------------------------------------------- এক্সেল


@login_required
def excel_page(request):
    return render(
        request,
        "core/excel.html",
        {"import_form": ExcelImportForm(), "today": local_today().isoformat()},
    )


@login_required
def excel_export(request):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Subjects"
    ws.append(["username", "display_name", "subject", "note", "submitted", "typed", "photocopied"])
    for s in Subject.objects.select_related("owner").order_by("owner__username", "order", "id"):
        ws.append(
            [
                s.owner.username,
                display_name(s.owner),
                s.name,
                s.note,
                "yes" if s.submitted else "no",
                "yes" if s.typed else "no",
                "yes" if s.photocopied else "no",
            ]
        )
    ws2 = wb.create_sheet("Todos")
    ws2.append(["username", "display_name", "text", "done"])
    for t in Todo.objects.select_related("owner").order_by("owner__username", "id"):
        ws2.append([t.owner.username, display_name(t.owner), t.text, "yes" if t.done else "no"])

    out = BytesIO()
    wb.save(out)
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="exam-section-{local_today().isoformat()}.xlsx"'
    )
    return response


@admin_required
def excel_template(request):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Subjects"
    ws.append(["username", "subject", "note", "submitted", "typed", "photocopied"])
    ws.append(["saimun", "বাংলা ১ম পত্র", "৩য় সেমিস্টার", "no", "no", "no"])
    ws2 = wb.create_sheet("Todos")
    ws2.append(["username", "text", "done"])
    ws2.append(["saimun", "প্রশ্নপত্র বন্ড করা", "no"])
    ws3 = wb.create_sheet("নির্দেশনা")
    for line in [
        "এক্সেল আমদানির নিয়ম:",
        "১) Subjects শিটে প্রতি সদস্যের বিষয়: username অবশ্যই অ্যাপে থাকা ইউজারনেম হতে হবে।",
        "২) একই ইউজারনেম + বিষয়ের নাম আগে থাকলে হালনাগাদ হবে, না থাকলে নতুন যোগ হবে।",
        "৩) submitted / typed / photocopied ঘরে yes বা no লিখুন (খালি রাখলে আগের মান থাকে)।",
        "৪) Todos শিটে: username, text, done (yes/no)। একই লেখা আগে থাকলে নতুন করে যোগ হবে না।",
        "৫) অতিরিক্ত কলাম থাকলে উপেক্ষা করা হয়; খালি সারি উপেক্ষা করা হয়।",
    ]:
        ws3.append([line])

    out = BytesIO()
    wb.save(out)
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="exam-section-template.xlsx"'
    return response


def _parse_bool(value):
    """yes/no/হ্যাঁ/না ইত্যাদি থেকে True/False; খালি হলে None।"""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip()
    if not text:
        return None
    return text.lower() in {"1", "yes", "true", "y", "x", "done", "✓", "√"} or text in {
        "হ্যাঁ",
        "হা",
        "১",
    }


def _find_sheet(wb, name):
    return wb[name] if name in wb.sheetnames else None


@admin_required
@require_POST
def excel_import(request):
    from openpyxl import load_workbook

    form = ExcelImportForm(request.POST, request.FILES)
    if not form.is_valid():
        messages.error(request, "একটি .xlsx ফাইল দিন।")
        return redirect("excel_page")

    try:
        wb = load_workbook(form.cleaned_data["file"], data_only=True)
    except Exception:
        messages.error(request, "ফাইলটি পড়া যায়নি। সঠিক .xlsx ফাইল কিনা দেখুন।")
        return redirect("excel_page")

    created = updated = skipped = 0

    ws = _find_sheet(wb, "Subjects") or wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        messages.error(request, "ফাইলে কোনো সারি নেই।")
        return redirect("excel_page")
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    if "username" not in headers or "subject" not in headers:
        messages.error(
            request, "Subjects শিটে 'username' ও 'subject' কলাম নেই। টেমপ্লেট ডাউনলোড করে দেখুন।"
        )
        return redirect("excel_page")
    col = {name: headers.index(name) for name in headers if name}

    def cell(row, key):
        idx = col.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    for row in rows[1:]:
        username = cell(row, "username")
        name = cell(row, "subject")
        if username is None or name is None or not str(username).strip() or not str(name).strip():
            skipped += 1
            continue
        user = User.objects.filter(username__iexact=str(username).strip()).first()
        if user is None:
            skipped += 1
            continue
        name = str(name).strip()
        note = str(cell(row, "note") or "").strip()
        subject = Subject.objects.filter(owner=user, name__iexact=name).first()
        if subject:
            changed = False
            if note and subject.note != note:
                subject.note = note
                changed = True
            for field in STATUS_FIELDS:
                value = _parse_bool(cell(row, field))
                if value is not None and getattr(subject, field) != value:
                    setattr(subject, field, value)
                    changed = True
            if changed:
                subject.save()
                updated += 1
        else:
            order = (user.subjects.aggregate(m=Max("order"))["m"] or 0) + 1
            Subject.objects.create(
                owner=user,
                name=name,
                note=note,
                submitted=bool(_parse_bool(cell(row, "submitted"))),
                typed=bool(_parse_bool(cell(row, "typed"))),
                photocopied=bool(_parse_bool(cell(row, "photocopied"))),
                order=order,
            )
            created += 1

    ws_todos = _find_sheet(wb, "Todos")
    todos_created = 0
    if ws_todos is not None:
        trows = list(ws_todos.iter_rows(values_only=True))
        if trows:
            theaders = [str(h).strip().lower() if h is not None else "" for h in trows[0]]
            if "username" in theaders and "text" in theaders:
                tcol = {n: theaders.index(n) for n in theaders if n}
                for row in trows[1:]:
                    u_idx, t_idx = tcol.get("username"), tcol.get("text")
                    if u_idx is None or t_idx is None or u_idx >= len(row) or t_idx >= len(row):
                        continue
                    uname, text = row[u_idx], row[t_idx]
                    if uname is None or text is None or not str(text).strip():
                        continue
                    user = User.objects.filter(username__iexact=str(uname).strip()).first()
                    if user is None:
                        continue
                    text = str(text).strip()
                    d_idx = tcol.get("done")
                    done = bool(_parse_bool(row[d_idx])) if d_idx is not None and d_idx < len(row) else False
                    _, was_created = Todo.objects.get_or_create(owner=user, text=text, defaults={"done": done})
                    todos_created += 1 if was_created else 0

    log_activity(
        request.user,
        request.user,
        Activity.SUBJECT_ADD,
        f"এক্সেল আমদানি: {created}টি নতুন বিষয়, {updated}টি হালনাগাদ",
    )
    messages.success(
        request,
        f"আমদানি শেষ — নতুন বিষয় {to_bn(created)}টি, হালনাগাদ {to_bn(updated)}টি, "
        f"বাদ পড়েছে {to_bn(skipped)}টি সারি, নতুন করণীয় {to_bn(todos_created)}টি।",
    )
    return redirect("excel_page")


# ---------------------------------------------------------------- সেটিংস


@admin_required
def settings_page(request):
    if request.method == "POST":
        form = SettingsForm(request.POST)
        if form.is_valid():
            AppSetting.set(AppSetting.EXAM_TITLE, form.cleaned_data["exam_title"])
            messages.success(request, "সেটিংস সংরক্ষিত হয়েছে।")
            return redirect("settings_page")
    else:
        form = SettingsForm(initial={"exam_title": AppSetting.exam_title()})
    return render(request, "core/settings.html", {"form": form})
